from pptx import Presentation

# Create a presentation object
prs = Presentation()

# Title slide
slide_1 = prs.slides.add_slide(prs.slide_layouts[0])  # 0 is the layout for title slide
title = slide_1.shapes.title
subtitle = slide_1.placeholders[1]
title.text = "My Family"
subtitle.text = "A Fun Journey Through Family Members"

# Slide 2: Introduction to My Family
slide_2 = prs.slides.add_slide(prs.slide_layouts[1])  # 1 is for title and content
slide_2.shapes.title.text = "Introduction to My Family"
content = slide_2.placeholders[1]
content.text = "A family is a group of people related by blood or marriage. Let's meet my family!"

# Slide 3: Meet My Family Members
slide_3 = prs.slides.add_slide(prs.slide_layouts[1])
slide_3.shapes.title.text = "Meet My Family Members"
content = slide_3.placeholders[1]
content.text = "Mom: She cooks the best meals!\nDad: He loves to play games!\nSibling: My partner in adventures!"

# Slide 4: Family Fun Activities
slide_4 = prs.slides.add_slide(prs.slide_layouts[1])
slide_4.shapes.title.text = "Family Fun Activities"
content = slide_4.placeholders[1]
content.text = "1. Going to the park\n2. Movie nights\n3. Cooking together"

# Slide 5: Why Family is Important
slide_5 = prs.slides.add_slide(prs.slide_layouts[1])
slide_5.shapes.title.text = "Why Family is Important"
content = slide_5.placeholders[1]
content.text = "Families provide love, support, and fun! They are always there for us."

# Save the presentation
prs.save('My_Family_Presentation.pptx')
